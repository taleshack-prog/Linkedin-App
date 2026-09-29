"""Parser do export de analytics do LinkedIn (perfil pessoal, .xlsx).

O arquivo vem do botão "Exportar" da tela de análises do criador. Ele NÃO é
uma planilha bem-comportada:

  - Tudo é string. Números ("168") e datas ("23/09/2026") também.
  - As datas saem no locale do perfil: dd/mm/yyyy em pt-BR, mm/dd/yyyy em
    en-US. O formato é detectado varrendo o arquivo inteiro (ver _detectar_ordem).
  - Os nomes das abas saem no idioma do perfil, por isso a identificação é
    por ESTRUTURA (o que a aba contém), com o nome apenas como dica.
  - A aba de publicações tem DUAS tabelas lado a lado (uma ordenada por
    engajamento, outra por impressões), separadas por uma coluna vazia, e o
    cabeçalho está na linha 3 — a linha 1 é um aviso de limite.

A URL de cada publicação termina em "...-share-<id>-XXXX", e <id> é o mesmo
número do `linkedin_post_urn` que o Posthink já grava ao publicar. É esse
casamento que permite juntar MÉTRICA (daqui) com TEXTO (do nosso banco).
"""
from __future__ import annotations

import io
import re
from datetime import date

MAX_POSTS = 50          # teto que o próprio LinkedIn declara no arquivo
MAX_DEMO_ITENS = 12     # por dimensão demográfica

_RE_SHARE = re.compile(r"[-:](?:share|activity|ugcPost)[-:](\d{6,})", re.I)
_RE_SLUG = re.compile(r"/posts/[^/]*?_([a-z0-9\-]+?)-(?:share|activity|ugcPost)-\d", re.I)
_RE_DATA = re.compile(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{4})\s*$")
_RE_NUM = re.compile(r"-?[\d.,]+")


class ExportError(Exception):
    pass


# --------------------------------------------------------------------------
# conversões
# --------------------------------------------------------------------------
def _txt(v) -> str:
    return "" if v is None else str(v).strip()


def _num(v) -> int | None:
    """'1.234' / '168' / '53%' -> int. Devolve None se não houver número."""
    s = _txt(v).replace("%", "")
    if not s:
        return None
    m = _RE_NUM.search(s)
    if not m:
        return None
    limpo = m.group(0).replace(".", "").replace(",", ".")
    try:
        return int(round(float(limpo)))
    except ValueError:
        return None


def _detectar_ordem(valores: list[str]) -> str:
    """Decide entre 'dmy' e 'mdy' olhando TODAS as datas do arquivo.

    Um componente > 12 só pode ser dia, e isso resolve o arquivo inteiro.
    Sem evidência, assume dia-primeiro (o resto do mundo além dos EUA).
    """
    for s in valores:
        m = _RE_DATA.match(s)
        if not m:
            continue
        a, b = int(m.group(1)), int(m.group(2))
        if a > 12:
            return "dmy"
        if b > 12:
            return "mdy"
    return "dmy"


def _data(v, ordem: str) -> str | None:
    """'23/09/2026' -> '2026-09-23' (ISO). None se não parecer data."""
    m = _RE_DATA.match(_txt(v))
    if not m:
        return None
    a, b, ano = int(m.group(1)), int(m.group(2)), int(m.group(3))
    dia, mes = (a, b) if ordem == "dmy" else (b, a)
    try:
        return date(ano, mes, dia).isoformat()
    except ValueError:
        return None


def _share_id(url: str) -> str | None:
    m = _RE_SHARE.search(url or "")
    return m.group(1) if m else None


def _slug(url: str) -> str | None:
    """As hashtags que o LinkedIn embute na URL — dá o tema mesmo sem o texto."""
    m = _RE_SLUG.search(url or "")
    if not m:
        return None
    return m.group(1).replace("-", " ").strip() or None


def urn_para_share_id(urn: str | None) -> str | None:
    """'urn:li:share:7510032742705737728' -> '7510032742705737728'.

    É o que casa um Post do nosso banco com uma linha do export.
    """
    if not urn:
        return None
    m = re.search(r"(\d{6,})\s*$", urn.strip())
    return m.group(1) if m else None


# --------------------------------------------------------------------------
# leitura das abas
# --------------------------------------------------------------------------
def _linhas(ws) -> list[list[str]]:
    return [[_txt(c) for c in row] for row in ws.iter_rows(values_only=True)]


def _tem_url_de_post(linhas: list[list[str]]) -> bool:
    return any("linkedin.com/posts/" in c or "/feed/update/" in c for ln in linhas for c in ln)


def _serie_diaria(linhas: list[list[str]], ordem: str) -> list[dict]:
    """Aba com [data, impressões, engajamentos] por dia."""
    fora = []
    for ln in linhas:
        if not ln:
            continue
        d = _data(ln[0], ordem)
        if not d:
            continue
        nums = [_num(c) for c in ln[1:]]
        nums = [n for n in nums if n is not None]
        if not nums:
            continue
        fora.append({
            "data": d,
            "impressoes": nums[0],
            "engajamentos": nums[1] if len(nums) > 1 else None,
        })
    return fora


def _pares_rotulo_valor(linhas: list[list[str]], ordem: str) -> list[dict]:
    """Aba de demografia: [dimensão, valor, percentual]."""
    fora = []
    for ln in linhas:
        if len(ln) < 3:
            continue
        dim, val, pct = ln[0], ln[1], ln[2]
        if not dim or not val or "%" not in pct:
            continue
        p = _num(pct)
        if p is None:
            continue
        fora.append({"dimensao": dim, "valor": val, "percentual": p})
    return fora[:MAX_DEMO_ITENS * 4]


def _publicacoes(linhas: list[list[str]], ordem: str) -> list[dict]:
    """Funde as duas tabelas lado a lado, casando pelo share id.

    Localiza cada bloco pela coluna que contém URLs, e lê as 3 colunas a
    partir dela (URL, data, métrica). O nome da métrica sai do cabeçalho,
    que é a última linha antes dos dados — por isso é procurado, não fixado
    na linha 3.
    """
    col_urls: list[int] = []
    largura = max((len(ln) for ln in linhas), default=0)
    for c in range(largura):
        if any("linkedin.com/posts/" in ln[c] or "/feed/update/" in ln[c]
               for ln in linhas if c < len(ln)):
            col_urls.append(c)

    porid: dict[str, dict] = {}
    for c in col_urls:
        # cabeçalho do bloco: linha acima da primeira com URL nessa coluna
        prim = next((i for i, ln in enumerate(linhas)
                     if c < len(ln) and "linkedin.com" in ln[c]), None)
        if prim is None:
            continue
        rotulo = ""
        for i in range(prim - 1, -1, -1):
            cand = linhas[i][c + 2] if c + 2 < len(linhas[i]) else ""
            if cand:
                rotulo = cand.lower()
                break
        chave = "engajamentos" if ("engaj" in rotulo or "engag" in rotulo) else "impressoes"

        for ln in linhas[prim:]:
            url = ln[c] if c < len(ln) else ""
            if "linkedin.com" not in url:
                continue
            sid = _share_id(url)
            if not sid:
                continue
            reg = porid.setdefault(sid, {
                "share_id": sid, "url": url, "data": None,
                "impressoes": None, "engajamentos": None, "tema_url": _slug(url),
            })
            if reg["data"] is None and c + 1 < len(ln):
                reg["data"] = _data(ln[c + 1], ordem)
            if c + 2 < len(ln):
                v = _num(ln[c + 2])
                if v is not None:
                    reg[chave] = v

    posts = list(porid.values())
    posts.sort(key=lambda p: (p["impressoes"] or 0, p["engajamentos"] or 0), reverse=True)
    return posts[:MAX_POSTS]


# --------------------------------------------------------------------------
# entrada pública
# --------------------------------------------------------------------------
def parse_analytics_xlsx(data: bytes, filename: str = "") -> dict:
    """Lê o .xlsx e devolve um dicionário JSON-serializável.

    Campos ausentes vêm como None — o arquivo varia conforme o que a conta
    tem. Nunca inventa número: o que não veio, não vem.
    """
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ExportError("openpyxl não instalado no servidor") from exc

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExportError(f"Não foi possível abrir a planilha: {exc}") from exc

    abas = {nome: _linhas(wb[nome]) for nome in wb.sheetnames}
    if not abas:
        raise ExportError("Planilha sem abas")

    todas = [c for linhas in abas.values() for ln in linhas for c in ln]
    ordem = _detectar_ordem(todas)

    fora: dict = {
        "periodo": {"inicio": None, "fim": None, "dias": None},
        "resumo": {"impressoes": None, "alcance": None, "engajamentos": None,
                   "seguidores_total": None, "seguidores_novos": None},
        "serie_diaria": [], "publicacoes": [],
        "publico": [], "conteudo": [],
        "avisos": [],
    }

    # período: o nome do arquivo traz ISO e é a fonte menos ambígua
    m = re.search(r"(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})", filename or "")
    if m:
        fora["periodo"]["inicio"], fora["periodo"]["fim"] = m.group(1), m.group(2)

    demo_encontradas = 0
    for nome, linhas in abas.items():
        n = nome.lower()
        if _tem_url_de_post(linhas):
            fora["publicacoes"] = _publicacoes(linhas, ordem)
            continue

        serie = _serie_diaria(linhas, ordem)
        rotulos = " ".join(ln[0].lower() for ln in linhas if ln and ln[0])

        if serie and ("seguidor" in rotulos or "follower" in n or "seguidor" in n):
            fora["resumo"]["seguidores_novos"] = sum(
                s["impressoes"] for s in serie if s["impressoes"] is not None
            )
            for ln in linhas:
                if ln and ("seguidor" in ln[0].lower() or "follower" in ln[0].lower()):
                    fora["resumo"]["seguidores_total"] = _num(ln[1] if len(ln) > 1 else "")
                    break
            continue

        if serie:
            fora["serie_diaria"] = serie
            fora["resumo"]["impressoes"] = fora["resumo"]["impressoes"] or sum(
                s["impressoes"] for s in serie if s["impressoes"] is not None
            )
            eng = [s["engajamentos"] for s in serie if s["engajamentos"] is not None]
            if eng:
                fora["resumo"]["engajamentos"] = sum(eng)
            continue

        demo = _pares_rotulo_valor(linhas, ordem)
        if demo:
            # a 1ª aba de demografia é o público (seguidores); a 2ª, o alcance
            destino = "publico" if demo_encontradas == 0 else "conteudo"
            if "conteúdo" in n or "content" in n:
                destino = "conteudo"
            elif "público" in n or "audience" in n:
                destino = "publico"
            fora[destino] = demo
            demo_encontradas += 1
            continue

        # aba de totais (Descoberta): rótulo + número, sem série nem demografia
        for ln in linhas:
            if len(ln) < 2:
                continue
            rot, val = ln[0].lower(), _num(ln[1])
            if val is None:
                continue
            if "impress" in rot and fora["resumo"]["impressoes"] is None:
                fora["resumo"]["impressoes"] = val
            elif "alcanç" in rot or "alcanc" in rot or "reach" in rot:
                fora["resumo"]["alcance"] = val

    # período pelo conteúdo, quando o nome do arquivo não deu
    if not fora["periodo"]["inicio"] and fora["serie_diaria"]:
        datas = sorted(s["data"] for s in fora["serie_diaria"])
        fora["periodo"]["inicio"], fora["periodo"]["fim"] = datas[0], datas[-1]

    if fora["periodo"]["inicio"] and fora["periodo"]["fim"]:
        fora["periodo"]["dias"] = (
            date.fromisoformat(fora["periodo"]["fim"])
            - date.fromisoformat(fora["periodo"]["inicio"])
        ).days + 1

    if not fora["publicacoes"] and not fora["serie_diaria"]:
        raise ExportError(
            "Não encontrei dados de desempenho nesta planilha. "
            "Use o arquivo do botão 'Exportar' da sua tela de análises do LinkedIn."
        )

    # avisos honestos sobre o alcance do que foi lido
    d = fora["periodo"]["dias"]
    if d is not None and d <= 31:
        fora["avisos"].append(
            f"Período curto ({d} dias). O LinkedIn só permite períodos personalizados pelo "
            "aplicativo do celular; no computador ficam os períodos pré-definidos."
        )
    if fora["publicacoes"] and len(fora["publicacoes"]) >= MAX_POSTS:
        fora["avisos"].append(f"O LinkedIn limita o export a {MAX_POSTS} publicações.")

    return fora
